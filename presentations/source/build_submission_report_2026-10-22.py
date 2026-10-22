"""Build the one-page submission from frozen results and its editable Markdown.

First run --figure-only with the project analysis environment (matplotlib), then
run without flags with the Codex bundled Python (python-docx). No network or live
calculation is used. Render the DOCX separately with the document skill's
render_docx.py and inspect the resulting page before publishing.
"""
from pathlib import Path
import hashlib
import json
import re
import sys


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "docs/SUBMISSION_WRITEUP.md"
FIGURE = ROOT / "figures/submission_results_2026-10-22.png"
OUT = ROOT / "presentations/results_2026-10-22/interphaise_submission_2026-10-22.docx"
HAIKU = ROOT / "data/experiments/environment_v1/environment_summary.json"
DEEPSEEK = ROOT / "data/experiments/deepseek_v1/environment_summary.json"


def figure():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.lines import Line2D
    from matplotlib.patches import Patch
    import numpy as np
    summaries = [json.loads(p.read_text()) for p in (HAIKU, DEEPSEEK)]
    environments = ["Li6PS5Cl", "Li3InCl6", "Li7La3Zr2O12"]
    titles = [r"Sulfide  Li$_6$PS$_5$Cl", r"Chloride  Li$_3$InCl$_6$", "Garnet  LLZO"]
    arms = ["iterative", "sequential", "shuffled", "upfront"]
    expected = [[[20, 2, 6, 0], [30, 30, 30, 30], [12, 0, 4, 1]],
                [[12, 2, 3, 2], [30, 30, 30, 30], [10, 4, 12, 20]]]
    colors = ["#007f85", "#cb7136"]
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 8,
                         "axes.titlesize": 9, "axes.labelsize": 8,
                         "svg.fonttype": "none"})
    fig, axes = plt.subplots(1, 3, figsize=(7.1, 1.83), sharey=True)
    y = np.arange(4)
    counts = {}
    for j, (ax, env, title) in enumerate(zip(axes, environments, titles)):
        counts[env] = {}
        for m, (summary, color) in enumerate(zip(summaries, colors)):
            rows = summary["environments"][env]["arms"]
            assert all(rows[a]["runs"] == 30 for a in arms)
            values = [rows[a]["acceptable"] for a in arms]
            assert values == expected[m][j], (env, m, values)
            counts[env][["Haiku", "DeepSeek"][m]] = values
            pos = y + (-0.17 if m == 0 else 0.17)
            ax.barh(pos, values, height=0.29, color=color, zorder=3)
            for py, value in zip(pos, values):
                ax.text(value + 0.5, py, str(value), va="center", fontsize=7.6,
                        color="#162c3e", zorder=5)
        baseline = summaries[0]["environments"][env]["arms"]["random"]["acceptable"]
        counts[env]["shared_random"] = baseline
        ax.axvline(baseline, color="#687482", ls=(0, (3, 2)), lw=1.2, zorder=4)
        ax.set_title(f"{title}\nRandom {baseline}/30", pad=6, fontsize=8.4)
        ax.set_xlim(0, 34)
        ax.set_xticks([0, 10, 20, 30])
        ax.set_yticks(y)
        ax.set_yticklabels(["Correct feedback", "No results", "Shuffled feedback", "Upfront shortlist"])
        ax.tick_params(axis="both", length=0, labelsize=8)
        ax.grid(axis="x", color="#e5e9ec", lw=0.6, zorder=0)
        for spine in ax.spines.values():
            spine.set_visible(False)
    axes[0].invert_yaxis()
    fig.legend(handles=[Patch(color=colors[0], label="Haiku 5.5"),
                        Patch(color=colors[1], label="DeepSeek V3.2"),
                        Line2D([0], [0], color="#687482", ls="--", label="Shared random")],
               loc="upper center", bbox_to_anchor=(0.58, 1.03), ncol=3,
               frameon=False, fontsize=8, handlelength=1.4, columnspacing=1.7)
    fig.text(0.59, 0.012, "Successful runs out of 30 within ten queries", ha="center", fontsize=8)
    fig.subplots_adjust(left=0.20, right=0.99, top=0.68, bottom=0.16, wspace=0.22)
    FIGURE.parent.mkdir(exist_ok=True)
    fig.savefig(FIGURE, dpi=300, facecolor="white")
    fig.savefig(FIGURE.with_suffix(".svg"), facecolor="white")
    svg = FIGURE.with_suffix(".svg")
    svg.write_text("\n".join(line.rstrip() for line in svg.read_text().splitlines()) + "\n")
    plt.close(fig)
    return counts


def add_link(paragraph, label, url, size):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    element = OxmlElement("w:hyperlink")
    element.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "006D77")
    props.append(color)
    font_size = OxmlElement("w:sz")
    font_size.set(qn("w:val"), str(round(size * 2)))
    props.append(font_size)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    element.append(run)
    paragraph._p.append(element)


def inline(paragraph, text, size=11):
    from docx.shared import Pt
    pos = 0
    for match in re.finditer(r"\[([^\]]+)\]\(([^)]+)\)", text):
        paragraph.add_run(text[pos:match.start()]).font.size = Pt(size)
        add_link(paragraph, match[1], match[2], size)
        pos = match.end()
    paragraph.add_run(text[pos:]).font.size = Pt(size)


def main():
    if sys.argv[1:] == ["--figure-only"]:
        counts = figure()
        FIGURE.with_suffix(".json").write_text(json.dumps(counts, indent=2) + "\n")
        print(FIGURE)
        return
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    counts = json.loads(FIGURE.with_suffix(".json").read_text())
    text = SOURCE.read_text()
    doc = Document()
    for border in doc.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    title = doc.styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(18)
    title.font.bold = True
    title.font.underline = False
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.keep_with_next = True

    for block in text.strip().split("\n\n"):
        if block.startswith("# "):
            doc.add_paragraph(block[2:], "Title")
        elif block.startswith("!["):
            match = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", block)
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = True
            pic = p.add_run().add_picture(str(SOURCE.parent / match[2]), width=Inches(7.1))
            pic._inline.docPr.set("descr", match[1])
        elif block.startswith("*") and block.endswith("*"):
            p = doc.add_paragraph()
            inline(p, block[1:-1], size=9)
            for run in p.runs:
                run.italic = True
        else:
            p = doc.add_paragraph()
            inline(p, block)
    doc.core_properties.title = "interphAIse for battery coating discovery"
    doc.core_properties.subject = "One-page OSI Hackathon submission report"
    doc.core_properties.author = "interphAIse team"
    doc.core_properties.comments = "Built from docs/SUBMISSION_WRITEUP.md and archived experiment summaries."
    doc.save(OUT)
    source_files = [SOURCE, HAIKU, DEEPSEEK, Path(__file__)]
    receipt = {"sources": {str(p.relative_to(ROOT)): hashlib.sha256(p.read_bytes()).hexdigest()
                            for p in source_files},
               "figure_counts": counts,
               "outputs": {str(p.relative_to(ROOT)): hashlib.sha256(p.read_bytes()).hexdigest()
                            for p in [OUT, FIGURE, FIGURE.with_suffix('.svg')]},
               "visual_validation": str(OUT.with_suffix('.validation.json').relative_to(ROOT))}
    OUT.with_suffix(".sources.json").write_text(json.dumps(receipt, indent=2) + "\n")
    print(OUT)
    print(FIGURE)


if __name__ == "__main__":
    main()
