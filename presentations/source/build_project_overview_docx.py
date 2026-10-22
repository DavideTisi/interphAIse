"""Build the approved short project overview with its graphical abstract.

Run with Python and python-docx from the Codex bundled document runtime.
The text and image are read from committed assets; no network access is used.
"""

from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "presentations/source/interphaise_overview_2026-10-22.txt"
IMAGE = ROOT / "presentations/results_2026-10-22/interphaise_graphical_abstract_2026-10-22_v1.png"
OUTPUT = ROOT / "presentations/results_2026-10-22/interphaise_project_overview_2026-10-22.docx"


def hyperlink(paragraph, label, url):
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "174F6B")
    props.append(color)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    props.append(size)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def main():
    doc = Document()
    # The runtime's default Word template contains accent rules on titles.
    for border in doc.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    section.header_distance = Inches(0.2)
    section.footer_distance = Inches(0.2)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.04
    normal.paragraph_format.widow_control = True
    for name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.underline = False
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.keep_with_next = True

    title = doc.add_paragraph("LLM Guided Battery Coating Discovery", "Title")
    title.runs[0].font.size = Pt(18)
    title.runs[0].bold = True
    meta = doc.add_paragraph("interphAIse  •  Project overview  •  22 October 2026")
    meta.runs[0].font.size = Pt(9.5)
    meta.paragraph_format.space_after = Pt(9)
    meta.paragraph_format.keep_with_next = True

    source_text = SOURCE.read_text(encoding="utf-8").strip()
    for index, text in enumerate(source_text.split("\n\n")):
        text = re.sub(r" \[Xiao et al\., Joule \(2019\)\]\([^)]*\)", " [1]", text)
        p = doc.add_paragraph()
        # Emphasize the project's question without changing the approved prose.
        if index == 1:
            first, rest = text.split(". ", 1)
            p.add_run(first + ". ").bold = True
            p.add_run(rest)
        else:
            p.add_run(text)

    figure = doc.add_paragraph()
    figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure.paragraph_format.space_before = Pt(2)
    figure.paragraph_format.space_after = Pt(2)
    figure.paragraph_format.keep_with_next = True
    picture = figure.add_run().add_picture(str(IMAGE), width=Inches(6))
    picture._inline.docPr.set("descr", "Conceptual summary of protective battery coatings, the LLM propose calculate interpret loop, and results. Sulfide: iterative 20 of 30 versus random 8 of 30. LLZO: iterative 12 of 30 versus random 24 of 30. LiTaF6 forecasts were in range; LiCaAlF6 forecasts missed. Atomic motifs are illustrative, not simulation outputs.")
    caption = doc.add_paragraph("Conceptual illustration. The atomic motifs are schematic, not the computed structures.")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].font.size = Pt(9)
    caption.paragraph_format.space_after = Pt(4)

    reference = doc.add_paragraph()
    reference.paragraph_format.space_after = Pt(0)
    reference.add_run("[1] Background: ").font.size = Pt(9)
    hyperlink(reference, "Xiao et al., Joule (2019)", "https://ceder.berkeley.edu/publications/Joule2019_coating.pdf")
    for r in reference.runs:
        r.font.size = Pt(9)

    doc.core_properties.title = "LLM Guided Battery Coating Discovery"
    doc.core_properties.subject = "interphAIse project overview and graphical abstract"
    doc.core_properties.author = "interphAIse team"
    doc.core_properties.keywords = "solid-state batteries, coatings, LLM, PET-MAD, scientific discovery"
    doc.core_properties.comments = "Built from the approved overview and illustration. Simulation evidence is provisional; no experimental battery validation is claimed."
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
