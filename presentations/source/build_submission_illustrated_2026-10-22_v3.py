"""Build the illustrated 810-run report using the bundled document runtime.

No scientific calculation or network access. Render separately and inspect every page.
The v2 figure is reused unchanged. Prose addresses the user-supplied judging rubric.
"""
from pathlib import Path
import hashlib
import json
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / 'presentations/source/interphaise_submission_illustrated_2026-10-22_v3.json'
FIGURE = ROOT / 'presentations/results_2026-10-22/interphaise_graphical_abstract_810_runs_2026-10-22_v2.png'
OUTPUT = ROOT / 'presentations/results_2026-10-22/interphaise_submission_illustrated_2026-10-22_v3.docx'
HAIKU = ROOT / 'data/experiments/environment_v1/environment_summary.json'
DEEPSEEK = ROOT / 'data/experiments/deepseek_v1/environment_summary.json'


def hyperlink(paragraph, label, url, size=9):
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    for tag, val in [('w:color', '006D77'), ('w:sz', str(round(size * 2)))]:
        element = OxmlElement(tag)
        element.set(qn('w:val'), val)
        props.append(element)
    run.append(props)
    text = OxmlElement('w:t')
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def verify_counts():
    envs = ['Li6PS5Cl', 'Li3InCl6', 'Li7La3Zr2O12']
    arms = ['iterative', 'sequential', 'shuffled', 'upfront']
    expected = [[[20, 2, 6, 0], [30, 30, 30, 30], [12, 0, 4, 1]],
                [[12, 2, 3, 2], [30, 30, 30, 30], [10, 4, 12, 20]]]
    summaries = [json.loads(p.read_text()) for p in (HAIKU, DEEPSEEK)]
    for mi, summary in enumerate(summaries):
        for ei, env in enumerate(envs):
            rows = summary['environments'][env]['arms']
            assert all(rows[a]['runs'] == 30 for a in arms)
            assert [rows[a]['acceptable'] for a in arms] == expected[mi][ei]
    random = [summaries[0]['environments'][e]['arms']['random']['acceptable'] for e in envs]
    assert random == [8, 29, 24]
    assert 2 * 4 * 3 * 30 + 3 * 30 == 810
    return {'arm_order': arms, 'environment_order': envs, 'Haiku': expected[0],
            'DeepSeek': expected[1], 'shared_random': random, 'distinct_runs': 810}


def main():
    counts = verify_counts()
    source = json.loads(SOURCE.read_text())
    doc = Document()
    for border in doc.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.55)
    section.left_margin = section.right_margin = Inches(0.7)
    section.header_distance = section.footer_distance = Inches(0.2)
    for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption'):
        style = doc.styles[name]
        style.font.name = 'Arial'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.underline = False
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.widow_control = True
    title = doc.styles['Title']
    title.font.size = Pt(18)
    title.font.bold = True
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.keep_with_next = True
    doc.add_paragraph(source['title'], 'Title')
    for index, block in enumerate(source['paragraphs']):
        p = doc.add_paragraph()
        p.add_run(block['lead'] + '. ').bold = True
        links = source.get('inline_links', {})
        pattern = '|'.join(re.escape(label) for label in links)
        pos = 0
        for match in re.finditer(pattern, block['text']):
            p.add_run(block['text'][pos:match.start()])
            hyperlink(p, match.group(), links[match.group()], 11)
            pos = match.end()
        p.add_run(block['text'][pos:])
        if index == 2:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = True
            shape = p.add_run().add_picture(str(FIGURE), width=Inches(7.1))
            shape._inline.docPr.set('title', 'Battery interface and LLM decisions with 810-run results')
            shape._inline.docPr.set('descr',
                'Schematic cathode, coating and electrolyte. The LLM proposes and interprets; '
                'code evaluates. Correct-feedback success out of 30 runs for Haiku and DeepSeek '
                'is 20 and 12 on sulfide, 30 and 30 on chloride, and 12 and 10 on LLZO. '
                'Shared random counts are 8, 29 and 24. In the separate Haiku pilot, both '
                'LiTaF6 adhesion forecasts held and both LiCaAlF6 forecasts missed.')
            p = doc.add_paragraph(source['caption'])
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.0
            p.runs[0].font.size = Pt(8.5)
            p.runs[0].italic = True

    p = section.footer.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for i, (label, url) in enumerate(source['links']):
        if i:
            p.add_run('  |  ').font.size = Pt(8.5)
        hyperlink(p, label, url, 8.5)
    p.add_run('  •  ' + source['footer']).font.size = Pt(8.5)
    doc.core_properties.title = source['title']
    doc.core_properties.subject = 'Illustrated one-page OSI submission addressing all four judging criteria'
    doc.core_properties.author = 'interphAIse team'
    doc.core_properties.keywords = 'solid-state batteries, coatings, LLM, feedback, PET-MAD, 810 runs'
    doc.core_properties.comments = 'Schematic generated illustration; numerical claims checked against archived evidence.'
    doc.save(OUTPUT)
    hashes = {str(p.relative_to(ROOT)): hashlib.sha256(p.read_bytes()).hexdigest()
              for p in [SOURCE, FIGURE, Path(__file__), HAIKU, DEEPSEEK, OUTPUT]}
    receipt = {'sha256': hashes, 'verified_counts': counts,
               'body_font_points': 11, 'page_format': 'Letter',
               'generation_method': 'Existing v2 image reused unchanged; no image generation in this revision',
               'visual_validation': 'Run render_docx.py and inspect before delivery'}
    OUTPUT.with_suffix('.sources.json').write_text(json.dumps(receipt, indent=2) + '\n')
    print(OUTPUT)


if __name__ == '__main__':
    main()
